
import os
import openpyxl
from pypdf import PdfReader

def read_excel(file_path):
    print(f"\n--- Reading Excel: {file_path} ---")
    try:
        wb = openpyxl.load_workbook(file_path)
        for sheet_name in wb.sheetnames:
            print(f"\nSheet: {sheet_name}")
            sheet = wb[sheet_name]
            for row in sheet.iter_rows(values_only=True):
                # Filter out None values to keep output clean
                cleaned_row = [str(cell) for cell in row if cell is not None]
                if cleaned_row:
                    print(cleaned_row)
    except Exception as e:
        print(f"Error reading Excel: {e}")

def read_pdf(file_path):
    print(f"\n--- Reading PDF: {file_path} ---")
    try:
        reader = PdfReader(file_path)
        for i, page in enumerate(reader.pages):
            print(f"\nPage {i+1}:")
            print(page.extract_text())
    except Exception as e:
        print(f"Error reading PDF: {e}")

if __name__ == "__main__":
    base_dir = "/Volumes/My Files/MathMax2K25/design"
    excel_path = os.path.join(base_dir, "MathMax_Levels.xlsx")
    pdf_path = os.path.join(base_dir, "MathMax.pdf")

    if os.path.exists(excel_path):
        read_excel(excel_path)
    else:
        print(f"File not found: {excel_path}")

    if os.path.exists(pdf_path):
        read_pdf(pdf_path)
    else:
        print(f"File not found: {pdf_path}")
