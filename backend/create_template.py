import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Users'

# Headers
headers = ['first_name', 'last_name', 'email', 'role', 'grade_level', 'section', 'phone_number']
ws.append(headers)

# Style headers
for cell in ws[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    cell.alignment = Alignment(horizontal='center')

# Sample data
ws.append(['John', 'Doe', 'john.doe@example.com', 'student', '9', 'A', '1234567890'])
ws.append(['Jane', 'Smith', 'jane.smith@example.com', 'student', '9', 'B', '2345678901'])
ws.append(['Robert', 'Johnson', 'robert.j@example.com', 'teacher', '', '', '3456789012'])

# Adjust column widths
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))
    ws.column_dimensions[column].width = max_length + 2

# Save
wb.save('bulk_upload_template.xlsx')
print('Template created successfully!')
